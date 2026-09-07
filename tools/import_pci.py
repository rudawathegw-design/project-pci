"""Import the PCI DSS project plan + assessment gap report into one JSON model."""
import openpyxl, json, re, sys
from datetime import datetime, date

PLAN = r"C:/Users/RudawAbdulrahman/Downloads/PCI DSS  Project Plan for FIB (3).xlsx"
GAPS = r"C:/Users/RudawAbdulrahman/Downloads/Assessment Gap Report (9).xlsx"

def norm_date(v):
    if v is None: return ""
    if isinstance(v, (datetime, date)): return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    if not s: return ""
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%m/%d/%Y"):
        try: return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError: pass
    return s  # keep raw if unparseable

def clean(v):
    if v is None: return ""
    return re.sub(r"\s+", " ", str(v)).strip()

def norm_status(s):
    s = clean(s).lower()
    if not s: return ""
    if "complet" in s or s == "done": return "Completed"
    if "progress" in s: return "In Progress"
    if "delay" in s: return "Delayed"
    if "hold" in s: return "On Hold"
    if "not start" in s or "no start" in s: return "Not Started"
    return clean(s).title()

# ── Roadmap (project plan) ──────────────────────────────────────────────────
wb = openpyxl.load_workbook(PLAN, data_only=True)
ws = wb.active
rows = list(ws.iter_rows(values_only=True))
hdr = None
for i, r in enumerate(rows):
    if r and clean(r[0]) == "Milestone Title":
        hdr = i; break
phases, cur = [], None
for r in rows[hdr+1:]:
    r = list(r) + [None]*8
    mt, tt, sd, ed, ad, asg, st, cm = r[:8]
    if not any([clean(mt), clean(tt), clean(st)]): continue
    if clean(mt):
        cur = {"title": clean(mt), "tasks": []}
        phases.append(cur)
    if cur is None:
        cur = {"title": "General", "tasks": []}; phases.append(cur)
    if clean(tt):
        cur["tasks"].append({
            "title": clean(tt), "start": norm_date(sd), "end": norm_date(ed),
            "actual": norm_date(ad), "assignee": clean(asg),
            "status": norm_status(st), "comments": clean(cm),
        })
# phase rollup
for p in phases:
    ts = p["tasks"]; n = len(ts)
    done = sum(1 for t in ts if t["status"] == "Completed")
    p["total"] = n; p["done"] = done
    p["pct"] = round(100*done/n) if n else 0
    if n and done == n: p["status"] = "Completed"
    elif any(t["status"] in ("In Progress","Delayed") for t in ts): p["status"] = "In Progress"
    elif done: p["status"] = "In Progress"
    else: p["status"] = "Upcoming"

# ── Gaps (assessment) ───────────────────────────────────────────────────────
wbg = openpyxl.load_workbook(GAPS, data_only=True)
areas = []
summary_sheet = None
for ws in wbg.worksheets:
    if "summary" in ws.title.lower(): summary_sheet = ws; continue
    rows = list(ws.iter_rows(values_only=True))
    # find header row (has 'Sr. No.')
    hi = None
    for i, r in enumerate(rows):
        if r and any(clean(c).lower().startswith("sr") for c in r if c): hi = i; break
    if hi is None: continue
    hdrs = [clean(c).lower() for c in rows[hi]]
    def col(*names):
        for j, h in enumerate(hdrs):
            if any(nm in h for nm in names): return j
        return None
    ci = {
        "section": col("section"), "obs": col("observation"),
        "rec": col("recommendation"), "ev": col("evidence required","evidence req"),
        "status": col("status"), "assessor": col("assessor"),
        "client": col("client comment"), "fib": col("fib status"),
        "link": col("link"),
    }
    findings = []
    op = cl = 0
    for r in rows[hi+1:]:
        r = list(r) + [None]*15
        section = clean(r[ci["section"]]) if ci["section"] is not None else ""
        obs = clean(r[ci["obs"]]) if ci["obs"] is not None else ""
        if not section and not obs: continue
        stt = clean(r[ci["status"]]).lower() if ci["status"] is not None else ""
        st = "Closed" if "clos" in stt else ("Open" if "open" in stt else clean(r[ci["status"]] if ci["status"] is not None else ""))
        if st == "Closed": cl += 1
        elif st == "Open": op += 1
        link = clean(r[ci["link"]]) if ci["link"] is not None else ""
        jira = link if "atlassian" in link else ""
        findings.append({
            "section": section, "observation": obs,
            "recommendation": clean(r[ci["rec"]]) if ci["rec"] is not None else "",
            "evidence_required": clean(r[ci["ev"]]) if ci["ev"] is not None else "",
            "status": st, "fib_status": norm_status(r[ci["fib"]]) if ci["fib"] is not None else "",
            "assessor_comments": clean(r[ci["assessor"]]) if ci["assessor"] is not None else "",
            "client_comments": clean(r[ci["client"]]) if ci["client"] is not None else "",
            "jira": jira,
        })
    total = op + cl
    areas.append({"name": clean(ws.title), "open": op, "closed": cl, "total": total,
                  "pct": round(100*cl/total) if total else 0, "findings": findings})

g_open = sum(a["open"] for a in areas); g_closed = sum(a["closed"] for a in areas)
g_total = g_open + g_closed
model = {
    "project": "PCI DSS — First Iraq Bank",
    "jira_project": "FIBXPI",
    "roadmap": {"phases": phases},
    "gaps": {"areas": areas,
             "summary": {"open": g_open, "closed": g_closed, "total": g_total,
                         "pct": round(100*g_closed/g_total) if g_total else 0}},
}
out = sys.argv[1] if len(sys.argv) > 1 else "pci_data.json"
with open(out, "w", encoding="utf-8") as f:
    json.dump(model, f, ensure_ascii=False, indent=2)

# summary print
print(f"Phases: {len(phases)}  | roadmap tasks: {sum(len(p['tasks']) for p in phases)}")
for p in phases: print(f"   {p['pct']:3}%  {p['status']:12}  {p['title']}")
print(f"\nGap areas: {len(areas)}  | overall remediated: {model['gaps']['summary']['pct']}%  ({g_closed}/{g_total} closed)")
for a in areas: print(f"   {a['pct']:3}%  open {a['open']:2} / closed {a['closed']:2}  {a['name']}  (findings {len(a['findings'])})")
print(f"\nWrote {out}")
